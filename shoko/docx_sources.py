"""Read-only inputs for the approved DOCX/RTF operations."""
import re
from datetime import timedelta
from pathlib import Path
from openpyxl import load_workbook
from .adapters import read_excel, readable
from .docx_rules import action_counts, price_table, price_notes
from .docx_rules import fleet_by_model as fleet_day
from .storage import AppError


def expanded(path, moment):
    try:
        return path.format(year=moment.year, yy=f'{moment.year%100:02d}', month=f'{moment.month:02d}')
    except (KeyError, ValueError, IndexError):
        raise AppError('Неверный шаблон пути. Допустимы {year}, {yy} и {month}.') from None


def monthly_path(config, moment, processing):
    current = processing - timedelta(days=1)
    template = config.get('drivers', '')
    historical = (moment.year, moment.month) != (current.year, current.month)
    source = template if not historical or '{' in template else config.get('drivers_previous', '')
    expected = f'Начисление_{moment:%m-%y}.xlsx'
    if not source:
        raise AppError(f'Нет начислений за {moment:%d.%m.%Y}. Выберите файл «{expected}» в настройках файлов; операция не запущена.')
    path = expanded(source, moment)
    match = re.fullmatch(r'Начисление_(\d{2})-(\d{2})\.xlsx', Path(path).name, re.I)
    if match and (int(match[1]), int(match[2])) != (moment.month, moment.year % 100):
        raise AppError(f'Для {moment:%d.%m.%Y} нужен «{expected}», выбран «{Path(path).name}».')
    if historical and '{' not in template and Path(path).resolve() == Path(expanded(template, current)).resolve():
        raise AppError(f'Для {moment:%d.%m.%Y} нужен отдельный файл «{expected}» за предыдущий месяц.')
    try:
        return str(readable(path, {'.xlsx'}))
    except AppError as exc:
        raise AppError(f'Начисления за {moment:%d.%m.%Y}: {exc}') from None


def sheet_names(path):
    readable(path, {'.xlsx'})
    try:
        workbook = load_workbook(path, read_only=True)
        try:
            return workbook.sheetnames
        finally:
            workbook.close()
    except Exception:
        raise AppError(f'Не удалось открыть Excel «{Path(path).name}».') from None


def prices(config, processing):
    path = expanded(config.get('prices', ''), processing)
    if not path:
        raise AppError('Выберите «Приказ_цены_2026.xlsx» в настройках файлов.')
    name = sheet_names(path)[0]
    rows = read_excel(path, name, 10_000, 5)
    parsed = price_table(rows)
    columns = 'A/B/C — модель/год/признак, E → F, D → G' if parsed['layout'] == 'split' else 'строки 9–73, C → F, B → G'
    detail = f'Приказ цен: первый лист «{name}», {columns}; распознано моделей: {len(parsed["values"])}.'
    return rows, [path], ' '.join([detail] + price_notes(rows))


def actions(config, processing):
    moment = processing - timedelta(days=1)
    path = monthly_path(config, moment, processing)
    names = sheet_names(path)
    day = str(moment.day) if str(moment.day) in names else f'{moment.day:02d}'
    try:
        rows = read_excel(path, day, 100_000, 41, strict_shape=False)
        counts, known = action_counts(rows)
    except AppError as exc:
        raise AppError(f'Начисления за {moment:%d.%m.%Y}, файл «{Path(path).name}», лист «{day}»: {exc}') from None
    return rows, [path], (f'За {moment:%d.%m.%Y}: лист «{day}», AM/AO/H: колонки и данные проверены; '
                          f'строк по справочнику: {known}, с непустым H: {sum(counts.values())}.')


def fleet(config, processing):
    moments = [processing - timedelta(days=n) for n in (1, 8, 15)]
    matrices = {}; result = []; paths = []
    resolved = [monthly_path(config,moment,processing) for moment in moments]
    for moment,path in zip(moments,resolved):
        if path not in matrices:
            matrices[path] = read_excel(path, 'Кол_состав', 100_000, 291, strict_shape=False)
            paths.append(path)
        try:
            result.append(fleet_day(matrices[path], moment.day))
        except AppError as exc:
            raise AppError(f'Начисления за {moment:%d.%m.%Y}, «{Path(path).name}», лист «Кол_состав»: {exc}') from None
    detail = 'Кол_состав: проверены даты ' + ', '.join(m.strftime('%d.%m.%Y') for m in moments) + '; поиск модели и года по содержимому, без объединения обычных и каршеринговых машин.'
    return result, paths, detail
