"""Small artificial examples shipped with the app, never user/customer data."""
from pathlib import Path
from openpyxl import Workbook
from .calculation import TARGET
from .docx_rules import FLEET_MODEL_YEARS, FLEET_CAR_MAPPING


def create_samples(folder):
    folder = Path(folder); folder.mkdir(parents=True, exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = 'Кол_состав_08-26'
    ws.cell(1, 13, 6); ws.cell(1, 22, 13); ws.cell(1, 31, 19)
    ws.cell(2, 2, 'Volkswagen Polo Эконом')
    ws.cell(3, 2, 2019); ws.cell(3, 3, 1200); ws.cell(3, 4, 1300)
    for c in range(13, 39): ws.cell(3, c, c)
    wb.save(folder / 'prostoi.xlsx'); wb.close()
    wb = Workbook(); ws = wb.active; ws.title = '19'
    ws.cell(1, 1, 'Старое поле A'); ws.cell(1, 2, 'Старое поле B'); ws.cell(1, 23, 'Автомобиль')
    ws.cell(2, 1, 2019); ws.cell(2, 2, -100); ws.cell(2, 23, 'Volkswagen Polo №001')
    ws.cell(1, 8, 'Скидка/наценка'); ws.cell(1, 39, 'люкс'); ws.cell(1, 41, 'год ПТС')
    ws.cell(2, 8, -100); ws.cell(2, 39, 'Volkswagen Polo Эконом'); ws.cell(2, 41, 2019)
    fleet = wb.create_sheet('Кол_состав')
    fleet.cell(14,2,'Volkswagen Polo Эконом');fleet.cell(21,2,2019)
    for day in range(1,32):
        nv={19:20,12:28,5:19}.get(day,0)
        start=12+(day-1)*9
        for offset,title in enumerate(('всего','аренда','простой аренда','выкуп','выкуп простой','%','простой НВ','% НВ')):
            fleet.cell(2,start+offset,title)
        for offset,value in enumerate((31,5,2,24,0,2/7,nv,.25,0)):
            fleet.cell(21,start+offset,value)
    wb.save(folder / 'drivers.xlsx'); wb.close()
    wb=Workbook(); ws=wb.active; ws.title='Цены'
    ws.cell(9,1,'Volkswagen Polo 2019 г.'); ws.cell(9,2,1200); ws.cell(9,3,1300)
    wb.save(folder / 'prices.xlsx'); wb.close()
    wb = Workbook(); ws = wb.active; ws.title = TARGET
    ws.append(['Класс', 'Модель', 'Год']); ws.append(['Эконом', 'Volkswagen Polo', 2019])
    ws['F2'] = 999; ws['G2'] = 888; ws['H2'] = 777
    garage = wb.create_sheet('Гараж 20.08.2026')
    garage.append(['Класс', 'Модель', 'Год выпуска'] + [''] * 12 + [f'Конкурент {i}' for i in range(1, 41)])
    garage.append(['Эконом', 'Volkswagen Polo', 2019])
    for address, expression in {'L2':'=AVERAGE(P2:AY2)','N2':'=COUNTA(P2:BC2)','BE2':'=MIN(P2:BC2)','BF2':'=AVERAGE(BE2;BG2)','BG2':'=AVERAGE(P2:BC2)','BH2':'=AVERAGE(BI2;BG2)','BI2':'=MAX(P2:BC2)'}.items():
        garage[address]=expression
    for column, source in ((10,'N'),(11,'BE'),(12,'BF'),(13,'L'),(14,'BH'),(15,'BI')):
        ws.cell(2,column,f"='Гараж 20.08.2026'!{source}2")
    wb.save(folder / 'target.xlsx'); wb.close()
    (folder / 'garage.csv').write_text('Марка и модель;Тариф;Год выпуска;Цена аренды 7 / 0;Тип топлива;КПП;Таксопарк\nVolkswagen Polo;эконом;2019;1500;бензин;автомат;Тестовый парк\n', encoding='utf-8')
    return dict(prices=str(folder/'prices.xlsx'), drivers_previous='', prostoi=str(folder/'prostoi.xlsx'), drivers=str(folder/'drivers.xlsx'), csv=str(folder/'garage.csv'),
                target=str(folder/'target.xlsx'), target_mode='local', last_date='2026-08-20', test_copy_confirmed=True)
