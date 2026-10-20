"""Approved Garage repairs. Existing market ranges and pricing coefficients stay intact."""
import math
import re
from decimal import Decimal, ROUND_HALF_UP
from openpyxl.utils import column_index_from_string, get_column_letter
from .calculation import TARGET, cell, set_cell
from .docx_rules import model_key, target_rows
from .driver_prices import target_value_formula
from .garage import formula_text, ERROR_VALUES, summary_value
from .storage import AppError

MARKET = ('J','K','L','M','BE','BF','BG','BH','BI','BO','BQ')
MARKET_NUMBERS = ('BS','BT','BU','BV')
MARKET_DERIVED = ('CH','CI','CJ')
RECOMMENDATIONS = tuple(get_column_letter(c) for c in range(90,104))  # CL:CY
GUARD = re.compile(r'^=IF\(COUNT\(P(\d+):([A-Z]+)\1\)=0;"-";(.*)\)$',re.I|re.S)
READY_GUARD = re.compile(r'^=IF\(OR\(COUNT\(P(\d+):BC\1\)=0;COUNT\(\$BR\1;\$E\1;\$F\1;\$G\1(?:;\$BX\1)?\)<[45]\);"-";(.*)\)$',re.I|re.S)
COMMENT_GUARD = re.compile(r'^=IF\(COUNT\(P(\d+):BC\1\)=0;"Недостаточно рыночных данных";IF\(COUNT\(\$BR\1;\$E\1;\$F\1;\$G\1(?:;\$BX\1)?\)<[45];"Недостаточно исходных данных";(.*)\)\)$',re.I|re.S)


RENTAL_GUARD = re.compile(r'^=IF\(COUNT\(\$BX(\d+)\)=0;"-";(.*)\)$',re.I|re.S)
RENTAL_COMMENT_GUARD = re.compile(r'^=IF\(COUNT\(P(\d+):BC\1\)=0;"Недостаточно рыночных данных";IF\(\$BX\1="-";"Нет арендного парка";IF\(COUNT\(\$BR\1;\$E\1;\$F\1;\$G\1;\$BX\1\)<5;"Недостаточно исходных данных";(.*)\)\)\)$',re.I|re.S)


def rental_ratio_formula(row):
    rented=target_value_formula('Q')[1:]
    return (f'=LET(rented;{rented};IF(COUNT(rented;$E{row};$F{row})<3;"";'
            f'IF(rented+$E{row}=0;"-";$F{row}/(rented+$E{row}))))')


def rental_guard(text,row):
    return f'=IF(COUNT($BX{row})=0;"-";{unguard(text)[1:]})'


def rental_values(source,address):
    """Validate real fleet counts; total P is informational, rental Q+U is the base."""
    values={c:_number(cell(source,column_index_from_string(c)-1),f'{address}, {c}')
            for c in ('P','Q','U','W')}
    if any(v<0 for v in values.values()) or values['W']>values['U'] or values['Q']+values['U']>values['P']:
        raise AppError(f'{address}: некорректные количества машин P/Q/U/W; проверьте состав парка.')
    base=values['Q']+values['U']
    return values['P'],values['W']/base if base else '-'


def unguard(text):
    for pattern in (RENTAL_COMMENT_GUARD,READY_GUARD,COMMENT_GUARD,RENTAL_GUARD,GUARD):
        m=pattern.fullmatch(text)
        if m:return '='+m[m.lastindex]
    return text


def market_guard(text, row, end='BC', ready=False, comment=False):
    core=unguard(text)[1:]
    count=f'COUNT(P{row}:{end}{row})=0'
    inputs=f'COUNT($BR{row};$E{row};$F{row};$G{row};$BX{row})<5'
    if comment:
        return f'=IF({count};"Недостаточно рыночных данных";IF($BX{row}="-";"Нет арендного парка";IF({inputs};"Недостаточно исходных данных";{core})))'
    condition=f'OR({count};{inputs})' if ready else count
    return f'=IF({condition};"-";{core})'


def repair_garage_template(source, target, work_name):
    mapping=target_rows(target)
    raw=source.setdefault('raw',[row.copy() for row in source['values']])
    positions=source.setdefault('formula_cells',[[r,c] for r,row in enumerate(raw) for c,v in enumerate(row) if isinstance(v,str) and v.startswith('=')])
    checks=[]
    if any(isinstance(cell(row,90),str) and cell(row,90).startswith('=') for row in raw[1:]):
        title=str(cell(raw[0],90)).lower()
        if 'цель вверх' not in title and 'target_up' not in title:
            raise AppError('Гараж, CM1: неизвестная структура блока рекомендаций. Формулы не изменены.')
    for r,identity in enumerate(source['values'][1:],1):
        model,year=cell(identity,1),cell(identity,2)
        if not model and year in ('',None):continue
        key=model_key(model,year);n=r+1
        if key not in mapping:
            raise AppError(f'«{work_name}», строка {n}: не найдена модель «{model}», {year} в «{TARGET}».')
        original_d=str(cell(raw[r],3))
        # Legacy Простои C is the current-driver price (target G). Tenet's
        # explicitly approved link to F, and existing explicit F links, remain F.
        price_column='F' if key==model_key('Tenet T7',2026) or "!F" in original_d else 'G'
        if original_d.startswith('=LET(') and "!F:F" in original_d:price_column='F'
        changed_columns=[]
        def put(column,text):
            c=column_index_from_string(column)-1
            set_cell(raw,r,c,text)
            if [r,c] not in positions:positions.append([r,c])
            changed_columns.append(column)
        for column,origin in (('D',price_column),('E','U'),('F','W'),('G','X')):
            put(column,target_value_formula(origin))
        for column in MARKET:
            c=column_index_from_string(column)-1;text=cell(raw[r],c)
            if not isinstance(text,str) or not text.startswith('='):continue
            core=unguard(text)
            match=re.search(r'\b(?:MIN|MAX|AVERAGE)\(P'+str(n)+r':([A-Z]+)'+str(n)+r'\)',core,re.I)
            end=match[1].upper() if match else 'BC'
            if column in ('K','M'):
                parent=cell(raw[r],9 if column=='K' else 11)
                m=re.search(r'COUNT\(P'+str(n)+r':([A-Z]+)'+str(n)+r'\)',str(parent),re.I)
                if m:end=m[1].upper()
            put(column,market_guard(core,n,end))
        if isinstance(cell(raw[r],69),str) and cell(raw[r],69).startswith('='):
            put('BR',f'=IF($D{n}="";"";VALUE(SUBSTITUTE($D{n};" ";"")))')
        for column in MARKET_NUMBERS:
            text=cell(raw[r],column_index_from_string(column)-1)
            if not isinstance(text,str) or not text.startswith('='):continue
            core=unguard(text)
            masked=re.fullmatch(r'=IFERROR\((.*);\s*0\)',core,re.I|re.S)
            if masked:core='='+masked[1]
            put(column,market_guard(core,n))
        for column in MARKET_DERIVED:
            text=cell(raw[r],column_index_from_string(column)-1)
            if isinstance(text,str) and text.startswith('='):
                put(column,market_guard(text,n))
        rental_columns=[]
        for column,expression in (('BW',target_value_formula('P')),('BX',rental_ratio_formula(n))):
            text=cell(raw[r],column_index_from_string(column)-1)
            if isinstance(text,str) and text.startswith('='):
                rental_values(target[mapping[key]],f'{TARGET}, строка {mapping[key]+1}')
                put(column,expression)
                rental_columns.append(column)
        for column in ('CB','CK'):
            text=cell(raw[r],column_index_from_string(column)-1)
            if isinstance(text,str) and text.startswith('='):
                put(column,rental_guard(text,n))
        cm=cell(raw[r],90)
        if isinstance(cm,str) and cm.startswith('='):
            put('CM',f'=MIN(MAX($BU{n};$CI{n});$BV{n})')
        for column in RECOMMENDATIONS:
            text=cell(raw[r],column_index_from_string(column)-1)
            if isinstance(text,str) and text.startswith('='):
                put(column,market_guard(text,n,ready=True,comment=column in ('CX','CY')))
        checks.append(dict(sheet=work_name,row=n,model=model,year=year,price_column=price_column,
                           rental_columns=rental_columns,columns=sorted(set(changed_columns)),recommendations=[c for c in RECOMMENDATIONS if c in changed_columns]))
    return checks


def _number(value,address,blank=False):
    if blank and value in ('',None):return None
    if type(value) not in (int,float) or not math.isfinite(value):
        raise AppError(f'Ошибка расчёта: {address} — ожидалось число.')
    return value


def _same(actual,expected,address):
    numeric=type(actual) in (int,float) and type(expected) in (int,float)
    if not (math.isclose(actual,expected,rel_tol=1e-12,abs_tol=1e-8) if numeric else type(actual)==type(expected) and actual==expected):
        raise AppError(f'Ошибка расчётной сверки: {address}.')


def pricing_values(row):
    """Independent numeric oracle for the retained CL/CN/CO/CP arithmetic."""
    def v(column):return _number(cell(row,column_index_from_string(column)-1),column)
    price,low,mean,upper,maximum=(v(c) for c in ('BR','BS','BT','BU','BV'))
    safe,over,strain=(v(c) for c in ('CI','CJ','CK'))
    down=min(max(mean-.3*strain*v('CH')-.7*over-.05*v('CD')*mean,safe),upper)
    up=min(max(upper,safe),maximum)
    goal=down if strain>1 or over>0 else up if v('CG')>=v('CF') and strain<=.6 else min(max(upper,safe),maximum)
    rounded=float(Decimal(str((price+.5*(goal-price))/100)).quantize(Decimal('1'),rounding=ROUND_HALF_UP))*100
    active=min(maximum if maximum>0 else upper,max(rounded,math.ceil((safe+100)/100)*100))
    old=cell(row,67)
    cap=min(active,float(old) if type(old) in (int,float) else 9000000000)
    discount=200 if strain<=.5 else 300 if strain<=1 else 400 if strain<=1.5 else 500
    promo=min(cap,math.ceil(max(math.ceil(safe/100)*100,cap-500,cap-discount)/100)*100)
    return dict(CL=down,CM=up,CN=goal,CO=active,CP=promo,CQ=active-promo,CT=promo,CV=active)


def verify_garage_repairs(snapshot,plan):
    target=snapshot['sheets'][TARGET]['values'];mapping=target_rows(target)
    for check in plan.get('garage_checks',[]):
        sheet=check['sheet'];r=check['row']-1;row=snapshot['sheets'][sheet]['values'][r]
        key=model_key(check['model'],check['year'])
        if key not in mapping:raise AppError('После записи изменилась модель/год в «Предложении цен».')
        source=target[mapping[key]]
        for column,origin in (('D',check['price_column']),('E','U'),('F','W'),('G','X'),('BP','G')):
            expected=cell(source,column_index_from_string(origin)-1)
            _number(expected,f'{TARGET}, {origin}{mapping[key]+1}',blank=True)
            _same(cell(row,column_index_from_string(column)-1),expected,f'{sheet}, {column}{r+1}')
        for c,value in enumerate(row[:103]):
            if isinstance(value,str) and value in ERROR_VALUES:
                raise AppError(f'Ошибка формулы: «{sheet}», {get_column_letter(c+1)}{r+1}.')
        rental_columns=check.get('rental_columns',[])
        if rental_columns:
            total,ratio=rental_values(source,f'{TARGET}, строка {mapping[key]+1}')
            for column,expected in (('BW',total),('BX',ratio)):
                if column in rental_columns:
                    _same(cell(row,column_index_from_string(column)-1),expected,f'{sheet}, {column}{r+1}')
            for column in ('CB','CK'):
                if column not in check['columns']:continue
                if ratio=='-':expected='-'
                elif column=='CB':
                    goal=_number(cell(row,78),f'{sheet}, CA{r+1}')
                    expected=min(2.5,ratio/goal if goal>0 else 0)
                else:
                    expected=_number(cell(row,79),'CB')*(.6+.2*_number(cell(row,81),'CD')+.2*_number(cell(row,76),'BY'))
                _same(cell(row,column_index_from_string(column)-1),expected,f'{sheet}, {column}{r+1}')
        markets=[v for v in row[15:55] if type(v) in (int,float)]
        ready=all(type(cell(row,c)) in (int,float) for c in ((69,4,5,6,75) if rental_columns else (69,4,5,6)))
        for column in check['recommendations']:
            if not markets or not ready:
                reason='Недостаточно рыночных данных' if not markets else 'Нет арендного парка' if rental_columns and cell(row,75)=='-' else 'Недостаточно исходных данных'
                expected=reason if column in ('CX','CY') else '-'
                _same(cell(row,column_index_from_string(column)-1),expected,f'{sheet}, {column}{r+1}')
        if markets and ready and all(c in check['recommendations'] for c in ('CL','CM','CN','CO','CP')):
            for column,expected in pricing_values(row).items():
                if column in check['recommendations']:
                    _same(cell(row,column_index_from_string(column)-1),expected,f'{sheet}, {column}{r+1}')
        # Validate own market summaries, preserving each author's original range.
        planned=plan['after'][sheet]
        for column in ('J','L','N','BE','BF','BG','BH','BI'):
            c=column_index_from_string(column)
            if column not in check['columns'] and column!='N':continue
            value=cell(planned[r],c-1)
            if isinstance(value,dict):
                expected=summary_value(planned,r,c,native=True)
                _same(cell(row,c-1),'' if expected is None else expected,f'{sheet}, {column}{r+1}')
