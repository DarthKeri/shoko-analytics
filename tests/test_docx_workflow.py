"""Approved DOCX/RTF profile, synthetic data and archived pure-function oracle."""
import copy
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from shoko import docx_sources
from shoko.docx_rules import apply_prices, apply_actions, apply_fleet, price_values, price_table, price_notes, action_counts, fleet_day
from shoko.calculation import TARGET, GARAGE
from shoko.adapters import LocalTarget
from shoko.engine import Engine
from shoko.garage import prepare_garage
from shoko.storage import AppError, file_hash
from test_mvp import fixtures
from test_garage import snapshot, csv_rows


class RtfRulesTests(unittest.TestCase):
    def test_fleet_matches_93_original_rtf_reference_cases(self):
        fixture=json.loads((Path(__file__).parent/'fixtures/rtf_fleet_reference.json').read_text(encoding='utf-8'))
        rows=[['']*fixture['columns'] for _ in range(fixture['rows'])]
        for r,c,value in fixture['cells']: rows[r][c]=value
        for day, expected in fixture['expected'].items():
            with self.subTest(day=day):
                self.assertEqual(fleet_day(rows,int(day)),{tuple(item['key']):item['data'] for item in expected})

    def test_prices_read_C_to_F_B_to_G_and_leave_unmatched_rows(self):
        source=[[] for _ in range(73)]
        source[8]=['Volkswagen Polo 2019 г.',1200.9,1300.9]
        source[14]=['Volkswagen Polo 2019 г.',9000,9000]  # RTF skip row 15
        source[72]=['kia rio 2021 г.',1400,1500]
        target=[['Class','Model','Year'],['Эконом','Volkswagen Polo',2019,1,2,9,8,7],['Эконом','Unknown',2020,1,2,6,5,4]]
        after,count=apply_prices(source,target)
        self.assertEqual(count,1);self.assertEqual(after[1][5:8],[1300,1200,7]);self.assertEqual(after[2],target[2])
        self.assertEqual(price_values(source)[('Kia Rio',2021)],[1500,1400])
        self.assertEqual(target[1][5],9)

    @staticmethod
    def split_prices():
        rows=[[] for _ in range(8)]
        rows[7]=['Автомобиль','год выпуска','доп.признак','действующий водитель, руб.','новый водитель, руб.']
        rows += [['Volkswagen POLO ',2021,None,2100,1300],
                 ['Volkswagen POLO ',2018,' Каршеринг',1900,1100],
                 ['Kia OPTIMA ',2019,' Люкс',3700,3700],
                 ['Toyota CAMRY ',2021,' 2.5',3200,2700],
                 ['Toyota CAMRY ',2020,' Актив+',2900,2400],
                 ['Kia K5 ',2022,' Комфорт',4500,3200],
                 ['Kia RIO ',2021,None,2400,1500],  # Row 15 must not be skipped.
                 ['Nissan QASHQAI ',2019,None,2100,1500],
                 ['Tenet T77',2026,None,3500,3500],
                 ['Volkswagen POLO ',2019,' Каршеринг',2000,1200],
                 ['Skoda OCTAVIA',2019,' Каршеринг',1900,1300],
                 ['Kia OPTIMA',2018,None,3000,3000]]
        return rows

    def test_split_prices_match_variant_year_and_role_without_old_row_skips(self):
        parsed=price_table(self.split_prices())
        self.assertEqual(parsed['layout'],'split')
        self.assertEqual(parsed['values'],{
            ('Volkswagen Polo',2021):[1300,2100],('Volkswagen Polo K',2018):[1100,1900],
            ('Kia Optima Люкс',2019):[3700,3700],('Toyota Camry 2.5',2021):[2700,3200],
            ('Toyota Camry Актив+',2020):[2400,2900],('Kia K5 Комфорт',2022):[3200,4500],
            ('Kia Rio',2021):[1500,2400],('Nissan Qashqai',2019):[1500,2100],
            ('Kia Optima',2018):[3000,3000],('Skoda Octavia К',2019):[1300,1900],
            ('Volkswagen Polo K',2019):[1200,2000],('Tenet T7',2026):[3500,3500]})
        self.assertEqual(parsed['unmapped'],[])
        self.assertEqual(price_notes(self.split_prices()),[])
        self.assertNotIn(('Volkswagen Polo',2019),parsed['values'])
        self.assertNotIn(('Skoda Octavia',2019),parsed['values'])

    def test_split_prices_preserve_other_columns_and_unknown_target_models(self):
        target=[[],['Эконом','Volkswagen Polo',2021,1,2,3,4,5],['Эконом','Tenet T77',2026,1,2,7,8,9]]
        result,matched=apply_prices(self.split_prices(),target)
        self.assertEqual(matched,1);self.assertEqual(result[1][5:7],[1300,2100])
        self.assertEqual(result[1][:5],target[1][:5]);self.assertEqual(result[1][7],5)
        self.assertEqual(result[2],target[2])

    def test_two_approved_price_models_update_only_exact_existing_rows(self):
        target=[['Класс','Модель','Год'],
                ['Комфорт','Kia Optima',2018,10,20,'','',77],
                ['Комфорт','Kia Optima Люкс',2018,10,20,88,99,77],
                ['Эконом','Skoda Octavia К',2019,10,20,'','',66],
                ['Эконом','Skoda Octavia',2019,10,20,88,99,66],
                ['Эконом','Volkswagen Polo',2019,10,20,88,99,66],
                ['Комфорт +','Faw Bestune T77',2023,10,20,88,99,66]]
        before=copy.deepcopy(target)
        after,matched=apply_prices(self.split_prices(),target)
        self.assertEqual(matched,2)
        self.assertEqual(len(after),len(before))
        self.assertEqual(after[1][5:7],[3000,3000])
        self.assertEqual(after[3][5:7],[1300,1900])
        changed=[(r,c) for r,(old,new) in enumerate(zip(before,after)) for c,(a,b) in enumerate(zip(old,new)) if a!=b]
        self.assertEqual(changed,[(1,5),(1,6),(3,5),(3,6)])
        self.assertEqual(target,before)
        self.assertEqual(price_notes(self.split_prices()),[])

    def test_price_additions_do_not_create_missing_rows_or_change_legacy_dictionary(self):
        from shoko.docx_rules import PRICE_MODEL_MAPPING
        self.assertNotIn('kia optima 2018 г.',PRICE_MODEL_MAPPING)
        self.assertNotIn('skoda octavia 2019 г. (каршеринг)',PRICE_MODEL_MAPPING)
        target=[[],['Эконом','Volkswagen Polo',2021,1,2,0,0,5]]
        after,matched=apply_prices(self.split_prices(),target)
        self.assertEqual(matched,1);self.assertEqual(len(after),2)
        self.assertEqual(after[1][5:7],[1300,2100])

    def test_split_prices_validate_headers_before_reading_any_prices(self):
        for col,value in ((1,'не год'),(3,'новый водитель'),(4,'действующий водитель')):
            rows=self.split_prices();rows[7][col]=value
            with self.subTest(col=col),self.assertRaisesRegex(AppError,'A8:E8'):
                price_values(rows)

    def test_split_prices_reject_duplicate_keys_bad_years_and_missing_formula_cache(self):
        rows=self.split_prices();rows.append(rows[8].copy())
        with self.assertRaisesRegex(AppError,'повтор модели/года'):price_values(rows)
        for value in ('#Нет сохранённого результата формулы','не число',2021.5,21):
            rows=self.split_prices();rows[8][1]=value
            with self.subTest(year=value),self.assertRaises(AppError):price_values(rows)
        for col in (3,4):
            rows=self.split_prices();rows[8][col]='#Нет сохранённого результата формулы'
            with self.subTest(col=col),self.assertRaisesRegex(AppError,'результата формулы'):price_values(rows)

    def test_split_prices_scan_beyond_old_row_limit_and_preserve_zero_and_truncation(self):
        rows=self.split_prices()[:8] + [[] for _ in range(80)]
        rows.append(['Kia RIO',2021,None,0,1500.9])
        self.assertEqual(price_values(rows),{('Kia Rio',2021):[1500,0]})


    def test_actions_use_AM_AO_H_and_count_zero_minus_and_duplicate_rows(self):
        header=['']*41;header[7]='Скидка/наценка';header[38]='люкс';header[40]='год ПТС'
        rows=[header]
        for marker in (0,'-',100,-100,'',None):
            row=['']*41;row[7]=marker;row[22]='Unknown W';row[38]=' Volkswagen  POLO эконом ';row[40]=2019;rows.append(row)
        target=[[],['Эконом','Volkswagen Polo',2019,1,2,3,4,99],['Эконом','Kia Rio',2021,1,2,3,4,88]]
        after,count,notes=apply_actions(rows,target)
        self.assertEqual(count,4);self.assertEqual(after[1][7],4);self.assertEqual(after[2][7],0);self.assertEqual(notes,[])
        self.assertEqual(after[1][3:7],target[1][3:7])

    def test_missing_formula_cache_cannot_count_as_an_action(self):
        header=['']*41;header[7]='Скидка/наценка';header[40]='год ПТС'
        row=['']*41;row[38]='Volkswagen Polo Эконом';row[40]=2019
        for value in ('#Нет сохранённого результата формулы','#N/A'):
            row[7]=value
            with self.subTest(value=value),self.assertRaisesRegex(AppError,'формулы'):
                action_counts([header,row])

    def test_actions_no_filled_H_preserves_previous_values_like_rtf(self):
        header=['']*41;header[7]='Скидка/наценка';header[40]='год ПТС'
        row=['']*41;row[38]='Volkswagen Polo Эконом';row[40]=2019
        target=[[],['Эконом','Volkswagen Polo',2019,1,2,3,4,77]]
        after,count,notes=apply_actions([header,row],target)
        self.assertEqual(after,target);self.assertEqual(count,0);self.assertTrue(notes)

    def test_fleet_touches_only_rtf_columns_and_stores_numeric_percent(self):
        data={'total':10,'rent':7,'buyout':1,'buyout_idle':2,'idle':3,'nv':4,'percent_after':.125}
        target=[[],['Эконом','Volkswagen Polo',2019]+[99]*26,['Эконом','Unknown',2020]+[88]*26]
        after,count=apply_fleet({('Volkswagen Polo',2019):data},{('Volkswagen Polo',2019):data},{('Volkswagen Polo',2019):data},target)
        changed={i for i,(a,b) in enumerate(zip(after[1],target[1])) if a!=b}
        self.assertEqual(changed,{15,16,17,18,20,21,22,24,25,27,28})
        self.assertEqual(after[1][21],.3);self.assertEqual(after[1][25],.125);self.assertEqual(after[2],target[2])
        self.assertEqual(after[1][19],99);self.assertEqual(after[1][23],99);self.assertEqual(after[1][26],99)

    def test_empty_JO_cells_stay_empty_and_do_not_require_a_website_or_formula(self):
        before=snapshot()
        for sheet_key in ('values','raw'):
            for row in before['sheets'][TARGET][sheet_key][1:]:
                for col in range(9,15):row[col]=''
        original=copy.deepcopy(before['sheets'][TARGET]['values'])
        after={TARGET:copy.deepcopy(original)}
        info=prepare_garage(before,after,csv_rows(),date(2026,10,14),update_links=False)
        self.assertEqual(after[TARGET],original);self.assertEqual(info['garage']['empty_links'],12)
        self.assertEqual(info['garage']['links'],0);self.assertEqual(info['derived'],[])
        self.assertEqual(after['Гараж 14.10.2026'][1][15],1500)

    def test_existing_IF_and_IFERROR_links_are_preserved_and_evaluated(self):
        before=snapshot();name='Гараж 25.08.2026'
        raw=before['sheets'][TARGET]['raw']
        raw[1][9]=f'=IF(\'{name}\'!N2=""; "-"; \'{name}\'!N2)'
        raw[2][10]=f'=IF(\'{name}\'!BE3=""; "-"; \'{name}\'!BE3)'
        raw[2][12]=f'=IFERROR(\'{name}\'!L3;"-")'
        before['sheets'][name]['raw'][2][11]='=IFERROR(AVERAGE(P3:AY3);"-")'
        after={TARGET:copy.deepcopy(before['sheets'][TARGET]['values'])};original=copy.deepcopy(after[TARGET])
        info=prepare_garage(before,after,csv_rows(),date(2026,10,14),update_links=False)
        derived={(d['row'],d['col']):d['value'] for d in info['derived']}
        self.assertEqual(derived[(2,10)],1)
        self.assertEqual(derived[(3,11)],0)  # MIN(empty)=0, not an empty cell.
        self.assertEqual(derived[(3,13)],'-')  # AVERAGE(empty) is caught by IFERROR.
        self.assertEqual(derived[(3,12)],'#DIV/0!')
        self.assertEqual(after[TARGET],original)
        self.assertEqual(before['sheets'][TARGET]['raw'],raw)

    def test_garage_keeps_JO_and_requires_existing_valid_links(self):
        before=snapshot();after={TARGET:copy.deepcopy(before['sheets'][TARGET]['values'])};original=copy.deepcopy(after[TARGET])
        info=prepare_garage(before,after,csv_rows(),date(2026,10,14),update_links=False)
        self.assertEqual(after[TARGET],original);self.assertEqual(info['garage']['links'],12)
        before['sheets'][TARGET]['raw'][1][9]='bad'
        with self.assertRaisesRegex(AppError,'J:O'):
            prepare_garage(before,{TARGET:copy.deepcopy(original)},csv_rows(),date(2026,10,14),update_links=False)

    def test_dash_cells_and_custom_dash_formulas_are_preserved_without_prediction(self):
        before=snapshot();name='Гараж 25.08.2026'
        for c,text in ((9,'-'),(10,'="-"'),(11,'=CUSTOM("-")')):
            before['sheets'][TARGET]['raw'][1][c]=text
            before['sheets'][TARGET]['values'][1][c]='-'
        # Exact analogue of live O18: the dash formula intentionally uses BH.
        before['sheets'][TARGET]['raw'][1][14]='=IF(\'Гараж 25.08.2026\'!BH2="";"-";\'Гараж 25.08.2026\'!BH2)'
        after={TARGET:copy.deepcopy(before['sheets'][TARGET]['values'])}
        original=copy.deepcopy(after[TARGET])
        info=prepare_garage(before,after,csv_rows(),date(2026,10,14),update_links=False)
        self.assertEqual(after[TARGET],original)
        self.assertEqual({item['cell'] for item in info['garage']['skipped_links']},{'J2','K2','L2','O2'})
        self.assertFalse(any(item['row']==2 and item['col'] in (10,11,12,15) for item in info['derived']))
        kept={(item['row'],item['col']):item['value'] for item in info['preserved']}
        self.assertEqual(kept[(2,10)],'-')
        self.assertEqual(kept[(2,15)],'=IF(\'Гараж 14.10.2026\'!BH2="";"-";\'Гараж 14.10.2026\'!BH2)')

    def test_custom_dash_formula_remains_skipped_when_its_result_is_numeric(self):
        before=snapshot()
        value='=IFERROR(CUSTOM_FUNCTION();"-")'
        before['sheets'][TARGET]['raw'][1][14]=value
        before['sheets'][TARGET]['values'][1][14]=1234
        after={TARGET:copy.deepcopy(before['sheets'][TARGET]['values'])}
        info=prepare_garage(before,after,csv_rows(),date(2026,10,14),update_links=False)
        self.assertEqual([s['cell'] for s in info['garage']['skipped_links']],['O2'])
        self.assertEqual(after[TARGET][1][14],1234)
        self.assertTrue(any(e['row']==2 and e['col']==15 and e['value']==value for e in info['preserved']))

    def test_preserved_formula_like_text_is_not_changed_by_sheet_rename(self):
        before=snapshot();value='=IFERROR(CUSTOM_FUNCTION();"-")'
        # Store text that merely looks like a formula; it stays plain text.
        before['sheets'][TARGET]['raw'][1][14]=value
        before['sheets'][TARGET]['values'][1][14]=value
        before['sheets'][TARGET]['formula_cells']=[[r,c] for r,row in enumerate(before['sheets'][TARGET]['raw']) for c,v in enumerate(row) if isinstance(v,str) and v.startswith('=') and (r,c)!=(1,14)]
        after={TARGET:copy.deepcopy(before['sheets'][TARGET]['values'])}
        info=prepare_garage(before,after,csv_rows(),date(2026,10,14),update_links=False)
        self.assertTrue(any(e['row']==2 and e['col']==15 and e['value']==value for e in info['preserved']))

    def test_non_dash_bad_link_still_blocks_and_zero_matches_never_bypass(self):
        before=snapshot();name='Гараж 25.08.2026'
        before['sheets'][TARGET]['raw'][1][14]=f"='{name}'!BH2"
        with self.assertRaisesRegex(AppError,'другой Гараж'):
            prepare_garage(before,{TARGET:copy.deepcopy(before['sheets'][TARGET]['values'])},csv_rows(),date(2026,10,14),update_links=False)
        before['sheets'][TARGET]['raw'][1][14]='-'
        with self.assertRaisesRegex(AppError,'совпад'):
            prepare_garage(before,{TARGET:copy.deepcopy(before['sheets'][TARGET]['values'])},csv_rows(model='Unknown'),date(2026,10,14),update_links=False)

    def test_google_verification_checks_preserved_formula_text_even_if_result_is_skipped(self):
        from shoko.adapters import GoogleTarget
        before=snapshot();before['sheets'][TARGET]['raw'][1][14]='=IF(\'Гараж 25.08.2026\'!BH2="";"-";\'Гараж 25.08.2026\'!BH2)'
        after={TARGET:copy.deepcopy(before['sheets'][TARGET]['values'])}
        info=prepare_garage(before,after,csv_rows(),date(2026,8,25),update_links=False)
        from shoko.garage import comparison_before
        from shoko.adapters import changes_for
        plan=dict(before=before,after=after,**info);plan['changes']=changes_for(comparison_before(plan),after)
        result=copy.deepcopy(before)
        for name,rows in after.items():result['sheets'][name]['raw']=[[v['formula'] if isinstance(v,dict) else v for v in row] for row in rows]
        for entry in info['preserved']:result['sheets'][TARGET]['raw'][entry['row']-1][entry['col']-1]=entry['value']
        for entry in info['derived']:result['sheets'][TARGET]['values'][entry['row']-1][entry['col']-1]=entry['value']
        adapter=object.__new__(GoogleTarget)
        with patch.object(adapter,'snapshot',return_value=result):adapter.verify(plan)
        result['sheets'][TARGET]['raw'][1][14]='-'
        with patch.object(adapter,'snapshot',return_value=result),self.assertRaisesRegex(AppError,'сохранённая.*O2'):
            adapter.verify(plan)


class DocxEngineTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.paths=fixtures(self.root/'files');self.engine=Engine(self.root/'data')
        self.engine.save_settings(self.paths)
    def tearDown(self):
        self.engine.close();self.temp.cleanup()
    def wait(self):
        worker=self.engine.worker
        if worker:worker.join(20)
        self.assertFalse(self.engine.busy())
    def prepare(self,operation):
        self.engine.prepare('',self.paths['last_date'],operation=operation);self.wait()
        self.assertEqual(self.engine.state['phase'],'preview',self.engine.state['error'])
        return self.engine.plan
    def write(self,plan):
        self.engine.write(plan['id'],plan['checksum']);self.wait()
        self.assertEqual(self.engine.state['phase'],'done',self.engine.state['error'])

    def test_each_operation_requires_only_its_sources_and_preparation_never_writes(self):
        for operation,required,expected in (('prices',{'prices'}, {TARGET}),('actions',{'drivers'},{TARGET}),('fleet',{'drivers'},{TARGET}),('csv',{'csv'},{GARAGE})):
            with self.subTest(operation=operation):
                self.engine.save_settings({k:self.paths[k] if k in required else '' for k in ('prices','drivers','csv')})
                before=file_hash(self.paths['target']);plan=self.prepare(operation)
                self.assertEqual(set(plan['after']),expected);self.assertEqual(plan['selected_operation'],operation)
                self.assertEqual(file_hash(self.paths['target']),before);self.assertEqual(self.engine.store.copies(),[])
                self.assertEqual({Path(s['path']).name for s in plan['sources'].values()},{Path(self.paths[k]).name for k in required})

    def test_two_selected_operations_share_one_review_and_touch_only_their_columns(self):
        self.engine.save_settings({'csv':''})
        before=file_hash(self.paths['target'])
        self.engine.prepare('',self.paths['last_date'],selected_operations=['actions','prices']);self.wait()
        self.assertEqual(self.engine.state['phase'],'preview',self.engine.state['error'])
        plan=self.engine.plan
        self.assertEqual(plan['selected_operations'],['actions','prices'])
        self.assertEqual(plan['selected_operation'],'selected')
        self.assertEqual(set(plan['after']),{TARGET})
        self.assertEqual({c['col'] for c in plan['changes']},{6,7,8})
        self.assertEqual(file_hash(self.paths['target']),before);self.assertFalse(self.engine.store.copies())
        self.write(plan)
        self.assertEqual(len(self.engine.store.copies()),1)

    def test_two_sheets_csv_and_prices_need_no_driver_file(self):
        self.engine.save_settings({'drivers':''})
        self.engine.prepare('',self.paths['last_date'],selected_operations=['prices','csv']);self.wait()
        self.assertEqual(self.engine.state['phase'],'preview',self.engine.state['error'])
        self.assertEqual(self.engine.plan['selected_operations'],['csv','prices'])
        self.assertEqual(set(self.engine.plan['after']),{TARGET,GARAGE})
        self.assertFalse(self.engine.plan.get('garage'))

    def test_CSV_and_garage_selection_uses_new_file_in_same_plan(self):
        self.engine.save_settings({'prices':'','drivers':''})
        self.engine.prepare('',self.paths['last_date'],selected_operations=['garage','csv']);self.wait()
        self.assertEqual(self.engine.state['phase'],'preview',self.engine.state['error'])
        plan=self.engine.plan
        self.assertEqual(plan['selected_operations'],['csv','garage'])
        self.assertEqual(plan['after'][GARAGE][1][3],'1500')
        self.assertEqual(plan['after']['Гараж 20.08.2026'][1][15],1500)
        self.assertFalse(any(c['sheet']==TARGET for c in plan['changes']))

    def test_empty_or_invalid_selection_never_defaults_to_all(self):
        before=file_hash(self.paths['target'])
        for selection in ([],['unknown'],'prices',[{}],['all']):
            with self.subTest(selection=selection),self.assertRaises(AppError):
                self.engine.prepare('',self.paths['last_date'],selected_operations=selection)
        self.assertEqual(file_hash(self.paths['target']),before);self.assertIsNone(self.engine.plan)

    def test_select_all_matches_legacy_all_and_deduplicates_selection(self):
        first=self.prepare('all')
        expected=copy.deepcopy(first['changes'])
        self.engine.prepare('',self.paths['last_date'],selected_operations=['prices','fleet','actions','garage','csv']);self.wait()
        self.assertEqual(self.engine.plan['changes'],expected)
        self.assertEqual(self.engine.plan['selected_operations'],['csv','garage','actions','fleet','prices'])
        self.engine.prepare('',self.paths['last_date'],selected_operations=['prices','prices']);self.wait()
        self.assertEqual(self.engine.plan['selected_operations'],['prices'])

    def test_raw_CSV_then_Garage_can_run_without_any_configured_source_file(self):
        self.write(self.prepare('csv'))
        self.engine.save_settings({'csv':'','drivers':'','prices':''})
        plan=self.prepare('garage')
        self.assertFalse(plan['sources']);self.assertIn('Гараж 20.08.2026',plan['after'])
        self.write(plan)

    def test_all_updates_only_DOCX_target_columns_raw_csv_and_garage(self):
        plan=self.prepare('all')
        self.assertEqual(plan['rule'],'docx-rtf-v1');self.assertNotIn('Простои',plan['after']);self.assertNotIn('Начисления',plan['after'])
        self.assertFalse(any(c['sheet']==TARGET and 10<=c['col']<=15 for c in plan['changes']))
        self.assertEqual(plan['after'][TARGET][1][5:8],[1300,1200,1])
        self.write(plan)
        for _ in range(2):self.assertEqual(self.prepare('all')['changes'],[])
        with self.assertRaisesRegex(AppError,'Изменений нет'):self.engine.write(self.engine.plan['id'],self.engine.plan['checksum'])

    def test_current_price_file_status_preview_backup_write_and_repeat(self):
        from openpyxl import Workbook
        w=Workbook();w.active.title='цены_07.10.26'
        for row in RtfRulesTests.split_prices():w.active.append(row)
        w.active.append(['Volkswagen POLO',2019,None,2000,1200])
        # Deliberately conflicting history: only the first worksheet is read.
        old=w.create_sheet('цены_02.10.26');old.append(['ignored',99999])
        w.save(self.paths['prices']);w.close()
        source_hash=file_hash(self.paths['prices'])
        _,_,status=docx_sources.prices(self.engine.config(),date(2026,10,14))
        self.assertIn('E → F, D → G',status);self.assertIn('цены_07.10.26',status)
        self.assertNotIn('нет в согласованном справочнике',status)
        plan=self.prepare('prices')
        self.assertEqual({c['col'] for c in plan['changes']},{1,2,3,6,7,9,24})
        self.assertEqual(len(plan['model_rows']),2)
        self.assertFalse(plan['warnings'])
        self.assertFalse(self.engine.store.copies())
        self.write(plan)
        self.assertEqual(len(self.engine.store.copies()),1)
        self.assertEqual(self.prepare('prices')['changes'],[])
        self.assertEqual(file_hash(self.paths['prices']),source_hash)


    def test_broken_workbook_has_clear_source_error(self):
        Path(self.paths['prices']).write_text('not an Excel archive',encoding='utf-8')
        with self.assertRaisesRegex(AppError,'Не удалось открыть Excel.*prices.xlsx'):
            docx_sources.prices(self.engine.config(),date(2026,8,20))

    def test_selected_prices_source_is_protected_by_hash_and_backup_gate(self):
        plan=self.prepare('prices');before=file_hash(self.paths['target'])
        w=load_workbook(self.paths['prices']);w.active['C9']=1400;w.save(self.paths['prices']);w.close()
        self.engine.write(plan['id'],plan['checksum']);self.wait()
        self.assertEqual(self.engine.state['phase'],'error');self.assertEqual(file_hash(self.paths['target']),before);self.assertFalse(self.engine.store.copies())

    def test_source_changed_during_backup_blocks_before_target_write(self):
        plan=self.prepare('prices');before=file_hash(self.paths['target'])
        original=self.engine.store.backup_file
        def backup_then_change(*args,**kwargs):
            result=original(*args,**kwargs)
            wb=load_workbook(self.paths['prices']);wb.active['C9']=9999;wb.save(self.paths['prices']);wb.close()
            return result
        with patch.object(self.engine.store,'backup_file',side_effect=backup_then_change):
            self.engine.write(plan['id'],plan['checksum']);self.wait()
        self.assertEqual(self.engine.state['phase'],'error')
        self.assertIn('во время создания копии',self.engine.state['error'])
        self.assertEqual(file_hash(self.paths['target']),before)
        self.assertEqual(len(self.engine.store.copies()),1)

    def test_dates_across_months_use_previous_book_and_exact_days(self):
        previous=self.root/'files/Начисление_07-26.xlsx'
        w=load_workbook(self.paths['drivers']);w.save(previous);w.close()
        self.engine.save_settings({'drivers_previous':str(previous)})
        config=self.engine.config()
        seen=[]
        def reader(path,sheet,*a,**kw):
            seen.append((Path(path).name,sheet));return [['']*291 for _ in range(356)]
        with patch('shoko.docx_sources.read_excel',side_effect=reader),patch('shoko.docx_sources.fleet_day',side_effect=lambda rows,day:{'day':day}):
            values,paths,detail=docx_sources.fleet(config,date(2026,8,10))
        self.assertEqual(values,[{'day':9},{'day':2},{'day':26}]);self.assertEqual(seen,[('drivers.xlsx','Кол_состав'),('Начисление_07-26.xlsx','Кол_состав')])
        self.assertIn('26.07.2026',detail)

    def test_missing_previous_month_blocks_fleet_but_does_not_block_prices(self):
        with self.assertRaisesRegex(AppError,'26.07.2026.*Начисление_07-26'):
            docx_sources.fleet(self.engine.config(),date(2026,8,10))
        self.prepare('prices')

    def test_month_file_mismatch_or_same_historical_path_is_blocked(self):
        config=self.engine.config();config['drivers_previous']=config['drivers']
        with self.assertRaises(AppError):docx_sources.monthly_path(config,date(2026,7,26),date(2026,8,10))
        with self.assertRaisesRegex(AppError,'Начисление_07-26'):
            docx_sources.monthly_path(dict(config,drivers_previous=str(self.root/'Начисление_08-26.xlsx')),date(2026,7,26),date(2026,8,10))

    def test_zero_matches_blocks_all_before_backup(self):
        Path(self.paths['csv']).write_text(Path(self.paths['csv']).read_text(encoding='utf-8').replace('Volkswagen Polo','Toyota Vitz'),encoding='utf-8')
        before=file_hash(self.paths['target'])
        self.engine.prepare('',self.paths['last_date']);self.wait()
        self.assertEqual(self.engine.state['phase'],'error');self.assertIn('ни одного совпадения',self.engine.state['error'])
        self.assertEqual(file_hash(self.paths['target']),before);self.assertFalse(self.engine.store.copies())


if __name__=='__main__':unittest.main()
