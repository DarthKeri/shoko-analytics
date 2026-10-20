import copy
import json
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from shoko.adapters import GoogleTarget
from shoko.engine import Engine
from shoko.calculation import TARGET
from shoko.sheet_backups import planned_backups, backup_matrix, verify_matrix
from shoko.storage import AppError, file_hash
from test_garage import snapshot, plan_for
import test_garage
from test_mvp import fixtures


class SheetBackupTests(unittest.TestCase):
    def test_only_existing_modified_sheets_are_backed_up(self):
        before = snapshot()
        changed = [{'sheet':TARGET}, {'sheet':TARGET}, {'sheet':'New'}]
        copies = planned_backups(before,changed,[])
        self.assertEqual(copies,[{'source':TARGET,'target':'BF_'+TARGET}])
        self.assertEqual(planned_backups(before,[],[]),[])

    def test_names_never_overwrite_and_use_moscow_time(self):
        before=snapshot();before['metadata']['sheets'].append({'properties':{'title':('BF_'+TARGET).lower(),'sheetId':20}})
        moment=datetime(2026,10,15,22,3,4,tzinfo=timezone.utc)
        copies=planned_backups(before,[{'sheet':TARGET}],[],moment)
        self.assertEqual(copies[0]['target'],'BF_'+TARGET+'_20261016_010304')
        before['metadata']['sheets'].append({'properties':{'title':copies[0]['target'],'sheetId':21}})
        self.assertTrue(planned_backups(before,[{'sheet':TARGET}],[],moment)[0]['target'].endswith('_2'))

    def test_formula_dependents_are_copied_when_only_their_result_changes(self):
        before=snapshot()
        derived=[{'sheet':TARGET,'row':2,'col':10,'value':1500}]
        copies=planned_backups(before,[{'sheet':'Гараж 25.08.2026'}],[],derived=derived)
        self.assertIn({'source':TARGET,'target':'BF_'+TARGET},copies)
        derived[0]['value']=1000
        self.assertNotIn(TARGET,[item['source'] for item in planned_backups(before,[{'sheet':'Гараж 25.08.2026'}],[],derived=derived)])
        derived[0]['value']=1500
        self.assertEqual(planned_backups(before,[],[],derived=derived),[])

    def test_local_names_respect_excel_limit(self):
        before=snapshot();before['mode']='local';name='x'*31;before['sheets'][name]=before['sheets'][TARGET]
        result=planned_backups(before,[{'sheet':name}],[])
        self.assertLessEqual(len(result[0]['target']),31)

    def test_renamed_garage_copies_old_source_and_keeps_plan_compact(self):
        plan=plan_for();original=copy.deepcopy(plan)
        plan['sheet_backups']=planned_backups(plan['before'],plan['changes'],plan['operations'])
        self.assertIn('Гараж 25.08.2026',[b['source'] for b in plan['sheet_backups']])
        self.assertNotIn('Гараж 14.10.2026',[b['source'] for b in plan['sheet_backups']])
        self.assertEqual(plan['before'],original['before'])
        self.assertLess(len(json.dumps(plan))-len(json.dumps(original)),1000)

    def test_google_copies_precede_any_changes_and_keep_active_ids(self):
        plan=plan_for();plan['sheet_backups']=planned_backups(plan['before'],plan['changes'],plan['operations'])
        adapter=object.__new__(GoogleTarget);payload=json.loads(adapter.prepare_write(plan))
        requests=payload['requests'];new_ids=[]
        for backup in plan['sheet_backups']:
            req=next(r['duplicateSheet'] for r in requests if r.get('duplicateSheet',{}).get('newSheetName')==backup['target'])
            self.assertEqual(req['sourceSheetId'],plan['before']['sheets'][backup['source']]['id'])
            self.assertNotIn(req['newSheetId'],[7,9]);new_ids.append(req['newSheetId'])
        self.assertEqual(len(new_ids),len(set(new_ids)))
        first_live=next(i for i,r in enumerate(requests) if r.get('updateCells',{}).get('start',{}).get('sheetId') in (7,9))
        self.assertTrue(all(i<first_live for i,r in enumerate(requests) if r.get('duplicateSheet',{}).get('newSheetName','').startswith('BF_')))
        self.assertFalse(any(r.get('updateSheetProperties',{}).get('properties',{}).get('title','').startswith('BF_') for r in requests))
        self.assertFalse(any('deleteSheet' in r for r in requests))

    def test_google_copy_indices_follow_sequential_insertions_and_renames(self):
        plan = plan_for()
        # Actual order comes from properties.index, not metadata response order.
        meta = plan['before']['metadata']['sheets']
        meta[0]['properties']['index'] = 1
        meta[1]['properties']['index'] = 2
        meta.append({'properties': {'sheetId': 100, 'title': 'Cover', 'index': 0}})
        plan['sheet_backups'] = planned_backups(plan['before'], plan['changes'], plan['operations'])
        adapter = object.__new__(GoogleTarget)
        requests = json.loads(adapter.prepare_write(plan))['requests']
        order = [(100, 'Cover'), (7, 'Гараж 25.08.2026'), (9, TARGET)]
        copied = []
        for request in requests:
            if 'duplicateSheet' in request:
                item = request['duplicateSheet']
                source_index = next(i for i, row in enumerate(order) if row[0] == item['sourceSheetId'])
                self.assertEqual(item['insertSheetIndex'], source_index + 1)
                copied.append((item['sourceSheetId'], item['newSheetId']))
                order.insert(item['insertSheetIndex'], (item['newSheetId'], item['newSheetName']))
            elif 'updateSheetProperties' in request:
                props = request['updateSheetProperties']['properties']
                if 'title' in props:
                    index = next(i for i, row in enumerate(order) if row[0] == props['sheetId'])
                    order[index] = (props['sheetId'], props['title'])
        self.assertEqual(order[0], (100, 'Cover'))
        for source, duplicate in copied:
            start = next(i for i, row in enumerate(order) if row[0] == source)
            stop = next((i for i in range(start + 1, len(order)) if order[i][0] in (7, 9, 100)), len(order))
            self.assertIn(duplicate, [row[0] for row in order[start + 1:stop]])

    def test_google_freezes_formulas_but_not_literal_formula_like_text(self):
        source={'raw':[['=1+1','=literal','=IF(1;"";2)']], 'values':[[2,'=literal','']], 'formula_cells':[[0,0],[0,2]]}
        self.assertEqual(backup_matrix(source,'google'),[[2,'=literal','']])
        self.assertEqual(backup_matrix(source,'local'),source['raw'])
        verify_matrix([[2,'=literal']],[[2.0,'=literal','']],'BF')
        with self.assertRaisesRegex(AppError,'BF-копии'):verify_matrix([[3]],[[2]],'BF')

    def test_google_verifies_backup_content_and_detects_damage(self):
        plan=plan_for();plan['sheet_backups']=planned_backups(plan['before'],plan['changes'],plan['operations'])
        result=test_garage.GarageTests().verified_result(plan);copies={'sheets':{}}
        for n,item in enumerate(plan['sheet_backups'],50):
            source=plan['before']['sheets'][item['source']]
            copies['sheets'][item['target']]=dict(copy.deepcopy(source),id=n,raw=backup_matrix(source,'google'))
        adapter=object.__new__(GoogleTarget)
        with patch.object(adapter,'snapshot',side_effect=[copies,result]):adapter.verify(plan)
        copies['sheets'][plan['sheet_backups'][0]['target']]['raw'][0][0]='damaged'
        with patch.object(adapter,'snapshot',return_value=copies),self.assertRaisesRegex(AppError,'BF-копии'):adapter.verify(plan)

    def test_local_roundtrip_preserves_previous_cells_style_and_comments(self):
        from openpyxl.comments import Comment
        with tempfile.TemporaryDirectory() as temp:
            paths=fixtures(Path(temp)/'files');wb=load_workbook(paths['target']);wb[TARGET]['F2'].comment=Comment('Fixture note','Fixture');wb.save(paths['target']);wb.close()
            engine=Engine(Path(temp)/'data');engine.save_settings(paths);before=file_hash(paths['target'])
            try:
                engine.prepare('',paths['last_date']);engine.worker.join(20);plan=engine.plan
                self.assertEqual(engine.state['phase'],'preview',engine.state['error']);self.assertTrue(plan['sheet_backups'])
                self.assertEqual(file_hash(paths['target']),before);self.assertEqual(engine.store.copies(),[])
                engine.write(plan['id'],plan['checksum']);engine.worker.join(20)
                self.assertEqual(engine.state['phase'],'done',engine.state['error'])
                wb=load_workbook(paths['target'])
                try:
                    backup=wb['BF_'+TARGET];self.assertEqual(backup['F2'].value,999);self.assertEqual(backup['D2'].value,'=1+1')
                    self.assertEqual(backup['D2'].fill.fgColor.rgb,'00ABCDEF');self.assertEqual(backup['F2'].comment.text,'Fixture note')
                    self.assertEqual(backup['F2'].fill.fgColor.rgb,'00FFF2CC')
                    self.assertNotEqual(wb[TARGET]['F2'].fill.fgColor.rgb,'00FFF2CC')
                    self.assertEqual(wb.sheetnames.index('BF_'+TARGET), wb.sheetnames.index(TARGET)+1)
                    self.assertEqual(wb[TARGET]['F2'].value,1300);self.assertEqual(wb['Не изменять']['A1'].value,'=2+3')
                finally:wb.close()
                engine.prepare('',paths['last_date']);engine.worker.join(20)
                self.assertEqual(engine.plan['changes'],[]);self.assertEqual(engine.plan['sheet_backups'],[])
            finally:engine.close()

if __name__=='__main__':unittest.main()
