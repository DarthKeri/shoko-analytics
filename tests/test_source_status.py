"""Early validation of monthly idle data and manual/automatic Garage exports."""
import copy
import json
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch, Mock
from openpyxl import load_workbook
from shoko.engine import Engine
from shoko.adapters import GoogleTarget, LocalTarget
from shoko.storage import AppError, file_hash
from shoko.validation import read_prostoi_source, read_bot_source, check_bot_source
from shoko.garage import match_garage_rows
from test_mvp import fixtures
from test_garage import snapshot, csv_rows


class SourceStatusTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.paths = fixtures(Path(self.temp.name) / 'files')
        self.engine = Engine(Path(self.temp.name) / 'workspace')
        self.engine.save_settings(self.paths)

    def tearDown(self):
        self.engine.close()
        self.temp.cleanup()

    def check(self, name, processing='2026-08-20'):
        self.engine.check(name, processing)
        if self.engine.worker:
            self.engine.worker.join(15)
        self.assertFalse(self.engine.busy())
        return self.engine.checks[name]

    def edit_idle(self, edit):
        wb = load_workbook(self.paths['prostoi'])
        edit(wb)
        wb.save(self.paths['prostoi'])
        wb.close()

    def test_idle_success_includes_processing_date_month_and_required_days(self):
        result = self.check('prostoi')
        self.assertTrue(result['ok'], result['message'])
        for text in ('20.08.2026', 'Кол_состав_08-26', '19.08.2026', '13.08.2026', '06.08.2026', 'числовые данные'):
            self.assertIn(text, result['message'])

    def test_idle_missing_month_has_date_and_required_sheet(self):
        result = self.check('prostoi', '2026-09-20')
        self.assertFalse(result['ok'])
        self.assertIn('20.09.2026', result['message'])
        self.assertIn('Кол_состав_09-26', result['message'])

    def test_idle_first_half_requires_previous_month_and_handles_year_boundary(self):
        result = self.check('prostoi', '2026-08-02')
        self.assertFalse(result['ok'])
        self.assertIn('Кол_состав_07-26', result['message'])
        self.edit_idle(lambda wb: setattr(wb.active, 'title', 'Кол_состав_01-26'))
        result = self.check('prostoi', '2026-01-02')
        self.assertIn('Кол_состав_12-25', result['message'])

    def test_idle_valid_first_half_and_year_boundary_use_both_months(self):
        for processing, current, previous in [('2026-08-02', 'Кол_состав_08-26', 'Кол_состав_07-26'),
                                               ('2026-01-02', 'Кол_состав_01-26', 'Кол_состав_12-25')]:
            with self.subTest(processing=processing):
                def months(wb):
                    for other in wb.worksheets[1:]:
                        wb.remove(other)
                    ws=wb.active; ws.title=current
                    for c in range(1,346):
                        ws.cell(1,c).value=None
                    for c in range(3,346):
                        ws.cell(3,c,0)
                    old=wb.copy_worksheet(ws); old.title=previous
                    old.cell(1,157,19); old.cell(1,166,26)
                    ws.cell(1,13,1)
                self.edit_idle(months)
                result=self.check('prostoi',processing)
                self.assertTrue(result['ok'], result['message'])
                self.assertIn(previous,result['message'])

    def test_idle_missing_daily_header_is_not_available(self):
        self.edit_idle(lambda wb: setattr(wb.active['AE1'], 'value', None))
        result = self.check('prostoi')
        self.assertFalse(result['ok'])
        self.assertIn('19.08.2026', result['message'])

    def test_idle_missing_model_year_blocks_status(self):
        self.edit_idle(lambda wb: setattr(wb.active['B3'], 'value', 'not a year'))
        result = self.check('prostoi')
        self.assertFalse(result['ok'])
        self.assertIn('модель → год', result['message'])

    def test_idle_bad_base_price_is_checked_immediately(self):
        self.edit_idle(lambda wb: setattr(wb.active['C3'], 'value', 'bad-price'))
        result = self.check('prostoi')
        self.assertFalse(result['ok'])
        self.assertIn('не является числом', result['message'])

    def test_idle_bad_day_value_and_uncached_formula_are_rejected(self):
        for value in ('bad-day', '=1+1'):
            with self.subTest(value=value):
                self.edit_idle(lambda wb: setattr(wb.active['AF3'], 'value', value))
                self.assertFalse(self.check('prostoi')['ok'])

    def test_idle_zero_remains_valid(self):
        self.edit_idle(lambda wb: setattr(wb.active['C3'], 'value', 0))
        self.assertTrue(self.check('prostoi')['ok'])

    def test_csv_success_reports_schema_eligible_and_matched_rows(self):
        result = self.check('csv')
        self.assertTrue(result['ok'], result['message'])
        for text in ('garage.csv', 'строк: 1', 'подходящих предложений: 1', 'Гараж 20.08.2026»: 1'):
            self.assertIn(text, result['message'])

    def test_csv_missing_required_columns_blocks_before_any_google_connection(self):
        Path(self.paths['csv']).write_text('model;price\nPolo;1000\n', encoding='utf-8')
        self.engine.store.config['target_mode'] = 'google'
        with patch('shoko.engine.target') as adapter:
            result = self.check('csv')
            adapter.assert_not_called()
        self.assertFalse(result['ok'])
        self.assertIn('Марка и модель', result['message'])
        self.assertIn('garage.csv', result['message'])

    def test_csv_bad_price_year_or_no_eligible_offers(self):
        original = Path(self.paths['csv']).read_text(encoding='utf-8')
        for old, new, text in [('1500', 'broken', 'неверная цена'), ('2019', 'wrong', 'неверный год'),
                               ('бензин', 'газ', 'Нет подходящих предложений')]:
            with self.subTest(text=text):
                Path(self.paths['csv']).write_text(original.replace(old, new), encoding='utf-8')
                result = self.check('csv')
                self.assertFalse(result['ok'])
                self.assertIn(text, result['message'])

    def test_csv_zero_matches_is_reported_during_status(self):
        path = Path(self.paths['csv'])
        path.write_text(path.read_text(encoding='utf-8').replace('Volkswagen Polo', 'Toyota Roomy'), encoding='utf-8')
        result = self.check('csv')
        self.assertFalse(result['ok'])
        self.assertIn('ни одного совпадения', result['message'])

    def test_csv_duplicate_normalized_required_header_is_rejected(self):
        path = Path(self.paths['csv'])
        rows = [['Марка и модель', ' Марка и модель ', *csv_rows()[0][1:]],
                ['Volkswagen Polo', 'Volkswagen Polo', *csv_rows()[1][1:]]]
        path.write_text('\n'.join(';'.join(row) for row in rows), encoding='utf-8')
        result = self.check('csv')
        self.assertFalse(result['ok'])
        self.assertIn('повторяются', result['message'])

    def test_csv_missing_template_or_changed_competitor_columns(self):
        wb = load_workbook(self.paths['target'])
        del wb['Гараж 20.08.2026']
        wb.save(self.paths['target']); wb.close()
        result = self.check('csv')
        self.assertFalse(result['ok'])
        self.assertIn('нет листа', result['message'])

    def test_shared_matching_allows_partial_missing_prices(self):
        template = snapshot()['sheets']['Гараж 25.08.2026']['values']
        result = match_garage_rows(template, csv_rows())
        self.assertEqual((result['matched'], result['cleared']), (1, 1))
        self.assertEqual(result['prices'][2], [])

    def test_google_status_reads_only_garage_template_using_get(self):
        adapter = object.__new__(GoogleTarget)
        metadata = snapshot()['metadata']
        rows = snapshot()['sheets']['Гараж 25.08.2026']['values']
        with patch.object(adapter, 'check', return_value=metadata), patch.object(adapter, 'request', return_value={'valueRanges': [{'values': rows}]}) as request:
            name, returned = adapter.garage_template()
        self.assertEqual(name, 'Гараж 25.08.2026')
        self.assertEqual(returned, rows)
        self.assertEqual(request.call_args.args, ('GET', '/values:batchGet'))
        self.assertEqual(request.call_args.kwargs['params']['ranges'], "'Гараж 25.08.2026'!A:BC")

    def test_target_change_invalidates_csv_match_status(self):
        self.assertTrue(self.check('csv')['ok'])
        self.engine.save_settings({'target': '', 'test_copy_confirmed': False})
        self.assertNotIn('csv', self.engine.checks)

    def test_all_status_is_read_only_and_creates_no_backup_or_plan(self):
        before = {key: file_hash(self.paths[key]) for key in ('prostoi', 'drivers', 'csv', 'target')}
        with patch.object(LocalTarget, 'write') as write, patch.object(self.engine.store, 'backup_file') as backup:
            self.engine.check('all', '2026-08-20')
            if self.engine.worker:
                self.engine.worker.join(15)
            self.assertTrue(all(c['ok'] for c in self.engine.checks.values()))
            write.assert_not_called(); backup.assert_not_called()
        self.assertEqual(before, {key: file_hash(self.paths[key]) for key in before})
        self.assertEqual(self.engine.store.copies(), [])
        self.assertEqual(list((self.engine.store.root / 'plans').iterdir()), [])


if __name__ == '__main__':
    unittest.main()
