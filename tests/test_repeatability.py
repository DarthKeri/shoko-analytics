import copy
import tempfile
import unittest
from datetime import date
from pathlib import Path
from openpyxl import load_workbook
from shoko.adapters import equivalent, changes_for
from shoko.calculation import assemble_prostoi, format_source, PROSTOI
from shoko.engine import Engine
from shoko.storage import AppError, file_hash
from test_mvp import fixtures


def month_rows():
    rows = [[''] * 345 for _ in range(4)]
    rows[1][1] = 'Модель'
    rows[2][1] = 2019
    for c in range(12, 345):
        rows[1][c] = '%' if c % 9 == 4 else 'Значение'
        rows[2][c] = 0 if c % 2 else 0.125
    return rows


class RepeatabilityTests(unittest.TestCase):
    def test_new_header_is_used_on_first_run_not_old_target_header(self):
        current = month_rows()
        stale = copy.deepcopy(current)
        stale[1] = ['%'] * 345
        current_copy = copy.deepcopy(current)
        first = assemble_prostoi(current, [], [], date(2026, 10, 20))
        other = assemble_prostoi(current, [], stale, date(2026, 10, 20))
        self.assertEqual(first, other)
        self.assertEqual(first[2][13], '0.0%')
        self.assertEqual(first[2][14], 0.125)
        self.assertEqual(current, current_copy)

    def test_repeated_runs_are_stable_for_both_half_months_and_month_lengths(self):
        current, previous = month_rows(), month_rows()
        for processing in (date(2026, 10, 14), date(2026, 10, 15),
                           date(2026, 2, 2), date(2026, 3, 2),
                           date(2024, 3, 2), date(2026, 5, 2)):
            with self.subTest(date=processing):
                first = assemble_prostoi(current, previous, [], processing)
                second = assemble_prostoi(current, previous, first, processing)
                third = assemble_prostoi(current, previous, second, processing)
                self.assertEqual(first, second)
                self.assertEqual(second, third)
                self.assertEqual(changes_for({'sheets': {PROSTOI: {'values': first}}}, {PROSTOI: second}), [])

    def test_percent_marker_follows_previous_month_slice_to_destination(self):
        current, previous = month_rows(), month_rows()
        previous[1][147] = '%'
        previous[2][147] = 0.125
        first = assemble_prostoi(current, previous, [], date(2026, 10, 14))
        self.assertEqual(first[1][12], '%')
        self.assertEqual(first[2][12], '12.5%')
        self.assertEqual(first, assemble_prostoi(current, previous, first, date(2026, 10, 14)))

    def test_changed_source_value_is_detected_and_next_run_is_stable(self):
        current = month_rows()
        first = assemble_prostoi(current, [], [], date(2026, 10, 20))
        current[2][13] = .25
        updated = assemble_prostoi(current, [], first, date(2026, 10, 20))
        changes = changes_for({'sheets': {PROSTOI: {'values': first}}}, {PROSTOI: updated})
        self.assertEqual([(c['cell'], c['after']) for c in changes], [('N3', '25.0%')])
        self.assertEqual(updated, assemble_prostoi(current, [], updated, date(2026, 10, 20)))

    def test_removed_source_percent_header_does_not_keep_stale_target_type(self):
        current = month_rows()
        current[2][13] = .25
        first = assemble_prostoi(current, [], [], date(2026, 10, 20))
        self.assertEqual(first[2][13], '25.0%')
        current[1][13] = 'Количество'
        updated = assemble_prostoi(current, [], first, date(2026, 10, 20))
        self.assertEqual(updated[2][13], .25)
        self.assertEqual(updated, assemble_prostoi(current, [], updated, date(2026, 10, 20)))

    def test_unwritten_tail_and_formula_are_preserved_before_midmonth(self):
        current, previous = month_rows(), month_rows()
        existing = [[''] * 500 for _ in range(4)]
        existing[0][499] = '=SUM(A1:B1)'
        first = assemble_prostoi(current, previous, existing, date(2026, 10, 14))
        self.assertEqual(first[0][499], '=SUM(A1:B1)')
        self.assertEqual(first, assemble_prostoi(current, previous, first, date(2026, 10, 14)))

    def test_percent_representation_and_invalid_data_rules_are_preserved(self):
        self.assertEqual(format_source(0, 20, {20}, True), '0.0%')
        self.assertEqual(format_source(.125, 20, {20}, True), '12.5%')
        self.assertEqual(format_source('12.5%', 20, {20}, True), '12.5%')
        self.assertEqual(format_source(0, 21, {20}, True), 0)
        self.assertEqual(format_source('%', 20, {20}, False), '%')
        with self.assertRaises(AppError):
            format_source('broken', 20, {20}, True)

    def test_numeric_storage_types_do_not_create_changes_but_text_and_bool_do(self):
        for a, b in ((0, 0.0), (123, 123.0), (None, '')):
            self.assertTrue(equivalent(a, b))
        for a, b in ((0, False), (1, True), (0, '0'), (1, '01'), (1, 1.00001)):
            self.assertFalse(equivalent(a, b))
        before = {'sheets': {'S': {'values': [[0.0, 123.0, False, '01']]}}}
        self.assertEqual(changes_for(before, {'S': [[0, 123, False, '01']]}), [])

    def test_local_write_then_restart_and_two_preparations_have_no_changes(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            paths = fixtures(root / 'inputs')
            workbook = load_workbook(paths['prostoi'])
            workbook.active.cell(2, 40, '%')
            workbook.active.cell(3, 40, .125)
            workbook.save(paths['prostoi'])
            workbook.close()
            data = root / 'data'
            engine = Engine(data)
            try:
                engine.save_settings(paths)
                engine.prepare('https://garage.example.test/fleet/1', '2026-08-20', True)
                if engine.worker:
                    engine.worker.join(20)
                self.assertEqual(engine.state['phase'], 'preview', engine.state['error'])
                self.assertEqual(engine.plan['after'][PROSTOI][2][39], '12.5%')
                plan = engine.plan
                engine.write(plan['id'], plan['checksum'])
                if engine.worker:
                    engine.worker.join(20)
                self.assertEqual(engine.state['phase'], 'done', engine.state['error'])
                target_hash = file_hash(paths['target'])
                copies = len(engine.store.copies())
            finally:
                engine.close()
            for attempt in range(2):
                engine = Engine(data)
                try:
                    engine.prepare('https://garage.example.test/fleet/1', '2026-08-20', True)
                    if engine.worker:
                        engine.worker.join(20)
                    self.assertEqual(engine.state['phase'], 'preview', engine.state['error'])
                    self.assertEqual(engine.plan['changes'], [], engine.plan['changes'][:2])
                    self.assertEqual(len(engine.store.copies()), copies)
                    self.assertEqual(file_hash(paths['target']), target_hash)
                finally:
                    engine.close()
