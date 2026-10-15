import copy
import json
import tempfile
import unittest
import time
import urllib.request
import urllib.error
from datetime import date
from pathlib import Path
from unittest.mock import patch
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from shoko.engine import Engine
from shoko.adapters import LocalTarget, GoogleTarget, guard_target, read_csv, read_excel, LIVE_ID, google_value, changes_for
from shoko.calculation import number, assemble_prostoi, driver_counts, PROSTOI, TARGET, DRIVERS, GARAGE
from shoko.storage import AppError, file_hash, InstanceLock
from shoko.server import Server
from shoko.plans import load_plan, save_plan


def fixtures(folder):
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = 'Кол_состав_08-26'
    ws.cell(1, 13, 6); ws.cell(1, 22, 13); ws.cell(1, 31, 19)
    ws.cell(2, 2, 'Volkswagen Polo Эконом')
    ws.cell(3, 2, 2019); ws.cell(3, 3, 1200); ws.cell(3, 4, 1300)
    for c in range(13, 39): ws.cell(3, c, c)
    wb.save(folder / 'prostoi.xlsx'); wb.close()
    wb = Workbook(); ws = wb.active; ws.title = '19'
    ws.cell(1, 1, 'Год ПТС'); ws.cell(1, 2, 'Скидка/наценка'); ws.cell(1, 23, 'Автомобиль')
    ws.cell(2, 1, 2019); ws.cell(2, 2, -100); ws.cell(2, 23, 'Volkswagen Polo №001')
    wb.save(folder / 'drivers.xlsx'); wb.close()
    wb = Workbook(); ws = wb.active; ws.title = TARGET
    ws.append(['Класс', 'Модель', 'Год']); ws.append(['Эконом', 'Polo', 2019])
    ws['D2'] = '=1+1'; ws['D2'].fill = PatternFill('solid', fgColor='ABCDEF')
    ws['F2'] = 999; ws['G2'] = 888; ws['H2'] = 777
    untouched = wb.create_sheet('Не изменять'); untouched['A1'] = '=2+3'
    garage = wb.create_sheet('Гараж 20.08.2026')
    garage.append(['Класс', 'Модель', 'Год выпуска'] + [''] * 12 + [f'Конкурент {i}' for i in range(1, 41)] + [''] * 6)
    garage.append(['Эконом', 'Volkswagen Polo', 2019])
    for address, expression in {'L2': '=AVERAGE(P2:AY2)', 'N2': '=COUNTA(P2:BC2)', 'BE2': '=MIN(P2:BC2)', 'BF2': '=AVERAGE(BE2;BG2)', 'BG2': '=AVERAGE(P2:BC2)', 'BH2': '=AVERAGE(BI2;BG2)', 'BI2': '=MAX(P2:BC2)'}.items():
        garage[address] = expression
    wb.save(folder / 'target.xlsx'); wb.close()
    (folder / 'garage.csv').write_text('Марка и модель;Тариф;Год выпуска;Цена аренды 7 / 0;Тип топлива;КПП;Таксопарк\nVolkswagen Polo;эконом;2019;1500;бензин;автомат;Тестовый парк\n', encoding='utf-8')
    return dict(prostoi=str(folder/'prostoi.xlsx'), drivers=str(folder/'drivers.xlsx'), csv=str(folder/'garage.csv'), target=str(folder/'target.xlsx'), target_mode='local', last_date='2026-08-20', test_copy_confirmed=True)


class EngineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.paths = fixtures(Path(self.temp.name)/'files')
        self.engine = Engine(Path(self.temp.name)/'data')
        self.engine.save_settings(self.paths)
        self.original = file_hash(self.paths['target'])

    def tearDown(self):
        self.engine.close()
        for attempt in range(5):
            try:
                self.temp.cleanup(); break
            except OSError:
                if attempt == 4: raise
                time.sleep(0.1)

    def wait(self):
        thread = self.engine.worker
        if thread: thread.join(15)
        self.assertFalse(self.engine.busy(), 'Worker did not finish')

    def prepare(self):
        self.engine.prepare('https://garage.example.test/fleet/1','2026-08-20',True)
        self.wait(); self.assertEqual(self.engine.state['phase'], 'preview', self.engine.state['error'])
        return self.engine.plan

    def write(self, plan):
        self.engine.write(plan['id'],plan['checksum']); self.wait()

    def test_full_local_cycle_backup_and_single_use(self):
        plan = self.prepare()
        self.assertEqual(file_hash(self.paths['target']),self.original)
        self.assertEqual(plan['after'][TARGET][1][5:8], [1300,1200,1])
        self.assertEqual(self.engine.plan_page('full', TARGET)['data'][1][3], '=1+1')
        self.assertEqual(plan['after'][TARGET][1][15],31)
        self.assertEqual(plan['after'][TARGET][1][24],28)
        self.assertEqual(plan['after'][TARGET][1][27],19)
        self.write(plan)
        self.assertEqual(self.engine.state['phase'],'done',self.engine.state['error'])
        result = self.engine.public_state()['history'][0]
        manifest, _ = self.engine.store.copy(result['backup'])
        self.assertEqual(file_hash(Path(manifest['path'])/'original.xlsx'),self.original)
        wb = load_workbook(self.paths['target'])
        self.assertEqual(wb[TARGET]['F2'].value,1300)
        self.assertEqual(wb[TARGET]['D2'].value,'=1+1')
        self.assertEqual(wb[TARGET]['D2'].fill.fgColor.rgb,'00ABCDEF')
        self.assertEqual(wb['Не изменять']['A1'].value,'=2+3')
        self.assertEqual(set([PROSTOI,DRIVERS,TARGET,GARAGE])-set(wb.sheetnames),set())
        wb.close()
        with self.assertRaises(AppError): self.engine.write(plan['id'],plan['checksum'])

    def test_changed_source_blocks_write(self):
        plan = self.prepare()
        Path(self.paths['csv']).write_text('model;price\nPolo;1600\n',encoding='utf-8')
        self.write(plan)
        self.assertEqual(self.engine.state['phase'],'error')
        self.assertIn('Исходный файл изменился',self.engine.state['error'])
        self.assertEqual(file_hash(self.paths['target']),self.original)

    def test_changed_target_blocks_write(self):
        plan = self.prepare()
        wb=load_workbook(self.paths['target']); wb[TARGET]['E2']=55; wb.save(self.paths['target']); wb.close()
        changed = file_hash(self.paths['target'])
        self.write(plan)
        self.assertEqual(self.engine.state['phase'],'error')
        self.assertIn('Целевая таблица изменилась',self.engine.state['error'])
        self.assertEqual(file_hash(self.paths['target']),changed)

    def test_tampered_plan_blocks_write(self):
        plan = self.prepare()
        saved=load_plan(self.engine.store.root,plan['id']);saved['after'][TARGET][1][5]=666
        save_plan(self.engine.store.root,saved)
        self.write(plan)
        self.assertIn('план повреждён',self.engine.state['error'])
        self.assertEqual(file_hash(self.paths['target']),self.original)

    def test_backup_failure_prevents_any_write(self):
        plan = self.prepare()
        with patch.object(self.engine.store,'backup_file',side_effect=AppError('Копия недоступна')),patch.object(LocalTarget,'write') as writer:
            self.write(plan); writer.assert_not_called()
        self.assertEqual(file_hash(self.paths['target']),self.original)

    def test_verification_failure_records_uncertain_and_keeps_backup(self):
        plan=self.prepare()
        with patch.object(LocalTarget,'verify',side_effect=AppError('Ошибка сверки')): self.write(plan)
        history=self.engine.public_state()['history'][0]
        self.assertEqual(history['status'],'uncertain')
        self.engine.store.copy(history['backup'])

    def test_checks_are_read_only(self):
        self.engine.check('all','2026-08-20');self.wait()
        self.assertTrue(all(v['ok'] for v in self.engine.checks.values()))
        self.assertEqual(file_hash(self.paths['target']),self.original)

    def test_manual_csv_requires_explicit_link_confirmation(self):
        with self.assertRaises(AppError): self.engine.prepare('https://garage.example.test','2026-08-20')
        self.assertEqual(file_hash(self.paths['target']),self.original)

    def test_missing_file_stops_preparation(self):
        Path(self.paths['drivers']).unlink()
        self.engine.prepare('https://garage.example.test','2026-08-20',True);self.wait()
        self.assertEqual(self.engine.state['phase'],'error');self.assertIsNone(self.engine.plan)
        self.assertEqual(file_hash(self.paths['target']),self.original)

    def test_bad_number_never_becomes_zero(self):
        wb=load_workbook(self.paths['prostoi']);wb.active['D3']='oops';wb.save(self.paths['prostoi']);wb.close()
        self.engine.prepare('https://garage.example.test','2026-08-20',True);self.wait()
        self.assertIn('не является числом',self.engine.state['error'])
        self.assertEqual(file_hash(self.paths['target']),self.original)

    def test_formula_without_cache_blocks_calculated_input(self):
        wb=load_workbook(self.paths['prostoi']);wb.active['D3']='=1+1';wb.save(self.paths['prostoi']);wb.close()
        self.engine.prepare('https://garage.example.test','2026-08-20',True);self.wait()
        self.assertEqual(self.engine.state['phase'],'error');self.assertIn('формулы',self.engine.state['error'])

    def test_copy_permission_can_be_revoked(self):
        self.engine.save_settings({'test_copy_confirmed':False})
        self.assertFalse(self.engine.public_state()['target_allowed'])
        with self.assertRaises(AppError):self.engine.prepare('https://garage.example.test','2026-08-20',True)

    def test_saved_setup_does_not_require_wizard_on_restart(self):
        self.assertTrue(self.engine.config()['initialized'])
        self.engine.finish_setup();self.assertTrue(self.engine.config()['initialized'])

    def test_cancel_before_read_never_writes(self):
        with patch.object(self.engine,'check_cancel',side_effect=AppError('Подготовка отменена')):
            self.engine.prepare('https://garage.example.test','2026-08-20',True);self.wait()
        self.assertEqual(file_hash(self.paths['target']),self.original)


class RulesTests(unittest.TestCase):
    def test_live_id_blocked_even_when_allowlisted(self):
        with self.assertRaises(AppError):guard_target({'target_mode':'google','target':LIVE_ID,'allowed_targets':[LIVE_ID]})

    def test_test_copy_confirmation_required(self):
        with self.assertRaises(AppError):guard_target({'target_mode':'local','target':'C:/qa.xlsx','allowed_targets':[]})

    def test_source_cannot_be_target(self):
        path=str(Path('C:/source.xlsx').resolve())
        with self.assertRaises(AppError):guard_target({'target_mode':'local','target':path,'prostoi':path,'allowed_targets':[path]})

    def test_invalid_numeric_inputs(self):
        for value in ['broken',True,float('inf')]:
            with self.assertRaises(AppError):number(value,'test')
        self.assertEqual(number('12,5%','test'),.125)

    def test_text_headers_are_preserved_but_bad_data_is_blocked(self):
        rows = [['', 'Модель', '', '', 'Всего', 'Средняя цена', 'Процент'], ['', 2019, '', '', 12, 1234.6]]
        result = assemble_prostoi(rows, [], [], date(2026, 8, 20))
        self.assertEqual(result[0][4:7], ['Всего', 'Средняя цена', 'Процент'])
        self.assertEqual(result[1][4:6], ['12', '1235'])
        rows[1][4] = 'bad-number'
        with self.assertRaises(AppError):assemble_prostoi(rows, [], [], date(2026, 8, 20))

    def test_month_start_carries_previous_month_tail(self):
        current=[list(range(345))];previous=[list(range(1000,1345))]
        result=assemble_prostoi(current,previous,[],date(2026,8,2))
        self.assertEqual([int(v) for v in result[0][:12]],list(range(12)))
        self.assertEqual(result[0][12],1156)
        self.assertEqual(result[0][146],1290)
        self.assertEqual(result[0][147],12)
        self.assertEqual(result[0][425],290)

    def test_unambiguous_model_case_and_spacing_are_normalized(self):
        row = ['']*39; row[0] = 2019; row[1] = -100; row[22] = 'SKODA  OCTAVIA №001'
        self.assertEqual(driver_counts([['Год ПТС','Скидка/наценка'] + [''] * 20 + ['марка'],row]), {('Skoda Octavia',2019):1})

    def test_unknown_discounted_model_requires_mapping(self):
        row=['']*39;row[0]=2020;row[1]=-100;row[22]='Unknown'
        with self.assertRaises(AppError):driver_counts([['Год ПТС','Скидка/наценка'] + [''] * 20 + ['марка'],row])

    def test_csv_structure_errors(self):
        with tempfile.TemporaryDirectory() as folder:
            path=Path(folder)/'bad.csv'
            for contents in ['','a,a\n1,2','a;a\n1;2','a;b\n1;2;3']:
                path.write_text(contents,encoding='utf-8')
                with self.assertRaises(AppError):read_csv(path)

    def test_google_writes_raw_text_one_batch(self):
        adapter=object.__new__(GoogleTarget)
        plan={'before':{'sheets':{TARGET:{'id':4,'rows':50,'cols':30}},'metadata':{'sheets':[{'properties':{'sheetId':5,'title':'Other'}}]}},'after':{TARGET:[['=malicious()']],GARAGE:[['head'],['=evil()']]},'changes':[{'sheet':TARGET,'row':1,'col':1,'after':'=malicious()'},{'sheet':GARAGE,'row':2,'col':1,'after':'=evil()'}]}
        with patch.object(adapter,'request') as request:
            adapter.write(plan)
            request.assert_called_once()
            body=json.loads(request.call_args.kwargs['data'])
            self.assertEqual(body['requests'][0]['addSheet']['properties']['sheetId'],6)
            for request_item in body['requests'][1:]:
                self.assertEqual(request_item['updateCells']['fields'],'userEnteredValue')
                self.assertIn('stringValue',request_item['updateCells']['rows'][0]['values'][0]['userEnteredValue'])

    def test_instance_lock_prevents_duplicate_process(self):
        with tempfile.TemporaryDirectory() as folder:
            lock=InstanceLock(folder)
            try:
                with self.assertRaises(AppError):InstanceLock(folder)
            finally:lock.close()


class HttpTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.engine=Engine(Path(self.temp.name)/'data')
        web=Path(self.temp.name)/'web';web.mkdir();(web/'index.html').write_text('ok')
        self.server=Server(self.engine,web);self.server.start()

    def tearDown(self):
        self.server.close();self.engine.close();self.temp.cleanup()

    def test_api_requires_token_and_same_origin(self):
        for headers in [{},{'X-Shoko-Token':self.server.token,'Origin':'https://untrusted.test'}]:
            with self.assertRaises(urllib.error.HTTPError) as ctx:
                urllib.request.urlopen(urllib.request.Request(self.server.url+'/api/state',headers=headers))
            self.assertEqual(ctx.exception.code,403)
        request=urllib.request.Request(self.server.url+'/api/state',headers={'X-Shoko-Token':self.server.token})
        self.assertIn('config',json.load(urllib.request.urlopen(request)))

    def test_bootstrap_rejects_cross_site_script(self):
        request=urllib.request.Request(self.server.url+'/bootstrap.js',headers={'Sec-Fetch-Site':'cross-site'})
        with self.assertRaises(urllib.error.HTTPError) as ctx:urllib.request.urlopen(request)
        self.assertEqual(ctx.exception.code,403)

    def test_get_action_cannot_write(self):
        request=urllib.request.Request(self.server.url+'/api/action',headers={'X-Shoko-Token':self.server.token})
        with self.assertRaises(urllib.error.HTTPError) as ctx:urllib.request.urlopen(request)
        self.assertEqual(ctx.exception.code,404)


if __name__=='__main__':unittest.main()
