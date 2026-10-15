"""Dated daily-source validation; synthetic XLSX only, no external writes."""
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from shoko.adapters import LocalTarget, read_driver_excel
from shoko.engine import Engine
from shoko.storage import AppError, file_hash
from test_mvp import fixtures


class DriverStatusTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.paths = fixtures(Path(self.temp.name) / 'files')
        self.engine = Engine(Path(self.temp.name) / 'data')
        self.engine.save_settings(self.paths)

    def tearDown(self):
        self.engine.close()
        self.temp.cleanup()

    def wait(self):
        worker = self.engine.worker
        if worker:
            worker.join(15)
        self.assertFalse(self.engine.busy())

    def edit(self, edit):
        wb = load_workbook(self.paths['drivers'])
        edit(wb)
        wb.save(self.paths['drivers'])
        wb.close()

    def check(self, processing='2026-08-20', source='drivers'):
        self.engine.check(source, processing)
        self.wait()
        return self.engine.checks['drivers']

    def test_success_reports_source_date_sheet_and_validated_columns(self):
        result = self.check()
        self.assertTrue(result['ok'], result['message'])
        self.assertIn('19.08.2026', result['message'])
        self.assertIn('лист «19»', result['message'])
        self.assertIn('колонки и данные проверены', result['message'])
        self.assertIn(result['message'], self.engine.store.logs())

    def test_missing_discount_blocks_individual_and_all_status(self):
        self.edit(lambda wb: setattr(wb['19']['B1'], 'value', None))
        for source in ('drivers', 'all'):
            with self.subTest(source=source):
                result = self.check(source=source)
                self.assertFalse(result['ok'])
                for value in ('19.08.2026', 'drivers.xlsx', 'лист «19»', 'Скидка/наценка', 'выгрузите начисления'):
                    self.assertIn(value, result['message'])
                self.assertIn(result['message'], self.engine.state['error'])

    def test_missing_sheet_explains_date_without_using_previous_sheet(self):
        result = self.check('2026-08-21')
        self.assertFalse(result['ok'])
        for value in ('20.08.2026', 'лист «20»', 'не найден лист'):
            self.assertIn(value, result['message'])

    def test_last_day_export_with_empty_required_area_is_not_available(self):
        def erase(wb):
            ws = wb['19']
            for row in ws.iter_rows(max_col=21):
                for cell in row:
                    cell.value = None
            # Preserve model/summary columns to reproduce the real old export.
            ws['AO1'] = 'год ПТС'
            ws['AO2'] = 2019
        self.edit(erase)
        result = self.check()
        self.assertFalse(result['ok'])
        self.assertIn('19.08.2026', result['message'])
        self.assertIn('Скидка/наценка', result['message'])

    def test_missing_year_never_falls_back_to_unrelated_am_column(self):
        self.edit(lambda wb: setattr(wb['19']['A1'], 'value', None))
        result = self.check()
        self.assertFalse(result['ok'])
        self.assertIn('Год ПТС', result['message'])

    def test_shifted_model_column_is_rejected(self):
        def shift(wb):
            wb['19']['X1'] = wb['19']['W1'].value
            wb['19']['X2'] = wb['19']['W2'].value
            wb['19']['W1'] = None
            wb['19']['W2'] = None
        self.edit(shift)
        result = self.check()
        self.assertFalse(result['ok'])
        self.assertIn('столбце W', result['message'])

    def test_duplicate_required_header_is_rejected(self):
        self.edit(lambda wb: setattr(wb['19']['C1'], 'value', 'Скидка/наценка'))
        result = self.check()
        self.assertFalse(result['ok'])
        self.assertIn('по одному разу', result['message'])

    def test_headers_without_records_are_not_available(self):
        def clear(wb):
            for row in wb['19'].iter_rows(min_row=2):
                for cell in row:
                    cell.value = None
        self.edit(clear)
        result = self.check()
        self.assertFalse(result['ok'])
        self.assertIn('нет строк', result['message'])

    def test_invalid_discounted_driver_is_reported_during_status(self):
        self.edit(lambda wb: setattr(wb['19']['A2'], 'value', None))
        result = self.check()
        self.assertFalse(result['ok'])
        self.assertIn('19.08.2026', result['message'])
        self.assertIn('Не указан год ПТС', result['message'])

    def test_zero_discount_is_valid_and_counts_as_populated(self):
        self.edit(lambda wb: setattr(wb['19']['B2'], 'value', 0))
        self.assertTrue(self.check()['ok'])

    def test_missing_file_error_includes_expected_date(self):
        self.engine.save_settings({'drivers': str(Path(self.temp.name) / 'absent.xlsx')})
        result = self.check()
        self.assertFalse(result['ok'])
        self.assertIn('19.08.2026', result['message'])
        self.assertIn('Файл не найден', result['message'])

    def test_month_and_year_boundary_resolves_previous_month_file_and_sheet(self):
        path = Path(self.temp.name) / 'Начисление_12-25.xlsx'
        wb = load_workbook(self.paths['drivers'])
        wb['19'].title = '31'
        wb.save(path)
        wb.close()
        self.engine.save_settings({'drivers': str(path.parent / 'Начисление_{month}-{yy}.xlsx')})
        result = self.check('2026-01-01')
        self.assertTrue(result['ok'], result['message'])
        self.assertIn('31.12.2025', result['message'])
        self.assertIn('лист «31»', result['message'])

    def test_source_and_target_unchanged_and_no_backup_or_plan_created(self):
        before = {k: file_hash(self.paths[k]) for k in ('drivers', 'prostoi', 'target', 'csv')}
        with patch.object(LocalTarget, 'write') as writer, patch.object(self.engine.store, 'backup_file') as backup:
            self.assertTrue(self.check(source='all')['ok'])
            writer.assert_not_called()
            backup.assert_not_called()
        self.assertEqual(before, {k: file_hash(self.paths[k]) for k in before})
        self.assertEqual(self.engine.store.copies(), [])
        self.assertEqual(list((self.engine.store.root / 'plans').iterdir()), [])

    def test_corrected_file_can_be_checked_again(self):
        self.edit(lambda wb: setattr(wb['19']['B1'], 'value', None))
        self.assertFalse(self.check()['ok'])
        self.edit(lambda wb: setattr(wb['19']['B1'], 'value', 'Скидка/наценка'))
        self.assertTrue(self.check()['ok'])
        self.assertEqual(self.engine.state['error'], '')

    def test_prepare_and_status_reject_the_same_invalid_source(self):
        self.edit(lambda wb: setattr(wb['19']['B1'], 'value', None))
        result = self.check()
        before = file_hash(self.paths['target'])
        self.engine.prepare('https://garage.example.test/fleet/1', '2026-08-20', True)
        self.wait()
        self.assertEqual(self.engine.state['error'], result['message'])
        self.assertIsNone(self.engine.plan)
        self.assertEqual(file_hash(self.paths['target']), before)
        self.assertEqual(self.engine.store.copies(), [])


if __name__ == '__main__':
    unittest.main()
